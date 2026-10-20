"""Render report/report.md as a Word document with the figures embedded.

    python -m scripts.build_docx

The Markdown file is the source of truth; this script only formats it. It
handles the subset of Markdown the report actually uses: headings, paragraphs
with inline bold/italic/code, bullet lists, tables, images with captions, and
horizontal rules.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path
from typing import List, Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from hw1.paths import ROOT

SOURCE = ROOT / "report" / "report.md"
OUTPUT = ROOT / "report" / "CS5370_Homework1_Report.docx"
PAGE_WIDTH_INCHES = 6.5

IMAGE = re.compile(r"^!\[(?P<alt>.*?)\]\((?P<path>.*?)\)\s*$")
CAPTION = re.compile(r"^\*(?P<text>Figure .+)\*\s*$")
INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)", re.DOTALL)


def add_runs(paragraph, text: str) -> None:
    """Write text into a paragraph, honouring **bold**, *italic* and `code`."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            paragraph.add_run(piece[1:-1]).italic = True
        else:
            paragraph.add_run(piece)


def split_row(line: str) -> List[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def add_table(document: Document, rows: List[List[str]]) -> None:
    header, body = rows[0], rows[1:]
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, header):
        cell.paragraphs[0].text = ""
        add_runs(cell.paragraphs[0], text)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for values in body:
        cells = table.add_row().cells
        for cell, text in zip(cells, values):
            cell.paragraphs[0].text = ""
            add_runs(cell.paragraphs[0], text)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    document.add_paragraph()


def add_image(document: Document, path: Path) -> None:
    if not path.exists():
        print(f"  missing figure: {path}")
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(PAGE_WIDTH_INCHES))


def add_caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    paragraph.paragraph_format.space_after = Pt(12)


def build() -> Path:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(8)
    for section in document.sections:
        section.left_margin = section.right_margin = Inches(1.0)
        section.top_margin = section.bottom_margin = Inches(0.9)

    index = 0
    pending_table: List[List[str]] = []

    def flush_table() -> None:
        nonlocal pending_table
        if pending_table:
            add_table(document, pending_table)
            pending_table = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("|"):
            # A separator row such as | --- | --- | carries no content.
            if not set(stripped) <= set("|-: "):
                pending_table.append(split_row(stripped))
            index += 1
            continue
        flush_table()

        if not stripped:
            index += 1
            continue

        if stripped == "---":
            index += 1
            continue

        image = IMAGE.match(stripped)
        if image:
            add_image(document, (SOURCE.parent / image.group("path")).resolve())
            index += 1
            continue

        caption = CAPTION.match(stripped)
        if caption:
            add_caption(document, caption.group("text"))
            index += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            heading = document.add_heading(level=min(level, 4))
            add_runs(heading, stripped[level:].strip())
            index += 1
            continue

        if stripped.startswith(("- ", "* ")):
            paragraph = document.add_paragraph(style="List Bullet")
            add_runs(paragraph, stripped[2:])
            index += 1
            continue

        if re.match(r"^\d+\.\s", stripped):
            paragraph = document.add_paragraph(style="List Number")
            add_runs(paragraph, re.sub(r"^\d+\.\s", "", stripped))
            index += 1
            continue

        # Gather a wrapped paragraph into one block before writing it.
        block = [stripped]
        index += 1
        while index < len(lines):
            following = lines[index].strip()
            if (
                not following
                or following.startswith(("#", "|", "- ", "* ", "!["))
                or following == "---"
                or CAPTION.match(following)
            ):
                break
            block.append(following)
            index += 1
        add_runs(document.add_paragraph(), " ".join(block))

    flush_table()
    document.save(OUTPUT)
    return OUTPUT


def main() -> int:
    path = build()
    print(f"wrote {path} ({path.stat().st_size / 1e6:.1f} MB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
